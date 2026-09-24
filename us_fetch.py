#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
us_fetch.py — 미국 지반(포지셔닝) 자료 수집기
=================================================================
왜 이 파일이 있는가
    조기경보 일일 회차의 L2(지반) 층은 「투자자가 빚을 얼마나 졌고 어느 쪽에
    쏠려 있나」를 잰다. 한국은 ews_fss.csv 가 그 자리를 채우지만 미국은
    2026-09-22 까지 «비어 있었다».
    원인은 클로드의 능력이 아니라 계정의 나가는 통신 허용 목록이었다 —
    2026-09-22 실측: 세션 프록시가 publicreporting.cftc.gov · cdn.finra.org ·
    cboe.com 에 대해 CONNECT 403(policy denial)을 돌려준다. 클라우드 컨테이너와
    사용자 PC 의 작업용 리눅스가 같은 목록을 쓰므로 양쪽 다 코드 000 이다.
    GitHub Actions 러너는 그 목록 밖에 있고, 러너가 커밋한 CSV 는
    raw.githubusercontent.com(허용됨)으로 세션이 읽는다 — 그래서 여기 있다.

무엇을 받는가 (셋 다 «원본 바이트»다. 요약 모델을 거치지 않는다)
    ① results/us/cftc_es.csv     선물 포지션  주 1회 (화요일 자료 · 금 15:30 ET 공개)
    ② results/us/finra_margin.csv 신용잔고     월 1회 (다음 달 셋째 주 공개)
    ③ results/us/cboe_putcall.csv 풋콜 비율    일 1회

설계 원칙 (krx_fetch.py · fss_fetch.py 와 같다)
    · **0행이면 쓰지 않는다.** 빈 응답이 좋은 행을 덮는 사고가 이미 났다(R-4·R-21).
    · 소스별로 독립 실패한다 — 하나가 죽어도 나머지는 수집하고 커밋한다.
    · 병합은 «날짜 키»로 한다. 재실행해도 같은 결과가 나온다(멱등).
    · 실패는 results/us/run_alarm.txt 로 빨간불이 되게 남긴다.
    · 매 실행의 소스별 결과를 results/_meta/us_source_health.csv 에 1행씩 적는다.
"""

import csv
import io
import json
import os
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from datetime import datetime, timezone

OUT_DIR = os.path.join("results", "us")
META_DIR = os.path.join("results", "_meta")
ALARM_PATH = os.path.join(OUT_DIR, "run_alarm.txt")
HEALTH_PATH = os.path.join(META_DIR, "us_source_health.csv")

RUN_UTC = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
UA = "Mozilla/5.0 (compatible; screen-us-fetch/1.0; +https://github.com/Leelll1/Screen)"

alarms = []
health_rows = []


def log(msg):
    print(msg, flush=True)


def http_get(url, timeout=60, retries=3, accept=None):
    """단순 GET. 실패하면 예외를 올린다 — 부르는 쪽이 소스별로 잡는다."""
    last = None
    for attempt in range(1, retries + 1):
        req = urllib.request.Request(url, headers={
            "User-Agent": UA,
            "Accept": accept or "*/*",
        })
        try:
            with urllib.request.urlopen(req, timeout=timeout) as r:
                return r.read()
        except Exception as e:  # noqa: BLE001 — 어떤 실패든 재시도 대상
            last = e
            log(f"    시도 {attempt}/{retries} 실패: {e}")
            if attempt < retries:
                time.sleep(attempt * 5)
    raise last


def merge_csv(path, fieldnames, new_rows, key):
    """
    기존 파일과 새 행을 키로 병합해 덮어쓴다.
    반환 — (전체 행수, 새로 생긴 행수, 값이 바뀐 행수)
    새 행이 0건이면 **아무것도 쓰지 않는다**(0행 덮어쓰기 방지).
    """
    if not new_rows:
        return (None, 0, 0)

    existing = {}
    if os.path.exists(path):
        with open(path, "r", encoding="utf-8", newline="") as f:
            for row in csv.DictReader(f):
                existing[row[key]] = row

    added = changed = 0
    for row in new_rows:
        k = row[key]
        if k not in existing:
            existing[k] = row
            added += 1
        else:
            old = existing[k]
            # 값이 실제로 달라진 경우에만 교체한다 — 재실행이 이력을 흔들지 않게.
            if any(str(old.get(c, "")) != str(row.get(c, "")) for c in fieldnames if c != "fetched_utc"):
                existing[k] = row
                changed += 1

    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
        w.writeheader()
        for k in sorted(existing):
            w.writerow(existing[k])
    return (len(existing), added, changed)


def record(source, status, detail, rows=None, added=None, changed=None):
    health_rows.append({
        "run_utc": RUN_UTC,
        "source": source,
        "status": status,
        "rows": "" if rows is None else rows,
        "added": "" if added is None else added,
        "changed": "" if changed is None else changed,
        "detail": detail,
    })
    log(f"  [{source}] {status} — {detail}")


# ─────────────────────────────────────────────────────────────────────
# ① CFTC — 선물 포지션 (주 1회)
# ─────────────────────────────────────────────────────────────────────
# 두 계열을 «둘 다» 받는다. 종전 회차들이 인용해 온 값은 legacy 계열이므로
# 계열을 바꾸면 이력이 끊긴다. TFF 는 헤지펀드/자산운용사를 갈라 주므로 더 낫다.
#   legacy  6dca-aqww  비상업(noncommercial) 롱/숏  ← 2026-09-19 회차까지 쓰던 값
#   TFF     gpe5-46if  딜러·자산운용사·레버리지드펀드
CFTC_SOURCES = [
    {
        "name": "cftc_legacy",
        "resource": "6dca-aqww",
        "name_field": "market_and_exchange_names",
        "like": "S&P 500 Consolidated%",
        "fields": [
            ("noncomm_positions_long_all", "noncomm_long"),
            ("noncomm_positions_short_all", "noncomm_short"),
            ("comm_positions_long_all", "comm_long"),
            ("comm_positions_short_all", "comm_short"),
            ("open_interest_all", "legacy_open_interest"),
        ],
    },
    {
        "name": "cftc_tff",
        "resource": "gpe5-46if",
        "name_field": "contract_market_name",
        "like": "E-MINI S&P 500%",
        "fields": [
            ("lev_money_positions_long", "lev_long"),
            ("lev_money_positions_short", "lev_short"),
            ("asset_mgr_positions_long", "am_long"),
            ("asset_mgr_positions_short", "am_short"),
            ("dealer_positions_long_all", "dealer_long"),
            ("dealer_positions_short_all", "dealer_short"),
            ("open_interest_all", "tff_open_interest"),
        ],
    },
]

CFTC_FIELDS = [
    "report_date", "contract",
    "noncomm_long", "noncomm_short", "noncomm_net",
    "comm_long", "comm_short", "legacy_open_interest",
    "lev_long", "lev_short", "lev_net",
    "am_long", "am_short", "am_net",
    "dealer_long", "dealer_short", "tff_open_interest",
    "fetched_utc",
]


def fetch_cftc(limit=700):
    """두 계열을 받아 report_date 로 합친다. limit=700 이면 주간 자료 약 13년치."""
    merged = {}
    got_any = False

    for src in CFTC_SOURCES:
        where = f"{src['name_field']} like '{src['like']}'"
        url = (
            f"https://publicreporting.cftc.gov/resource/{src['resource']}.json"
            f"?$where={urllib.parse.quote(where)}"
            f"&$order=report_date_as_yyyy_mm_dd%20DESC&$limit={limit}"
        )
        try:
            raw = http_get(url, accept="application/json")
            recs = json.loads(raw.decode("utf-8"))
        except Exception as e:  # noqa: BLE001
            record(src["name"], "FAIL", f"조회 실패: {type(e).__name__} {e}")
            alarms.append(f"{src['name']}: {e}")
            continue

        if not recs:
            record(src["name"], "EMPTY", "0행 — 쓰지 않는다")
            continue

        got_any = True
        for r in recs:
            d = (r.get("report_date_as_yyyy_mm_dd") or "")[:10]
            if not d:
                continue
            row = merged.setdefault(d, {"report_date": d, "fetched_utc": RUN_UTC})
            row.setdefault("contract", r.get(src["name_field"], ""))
            for api_name, our_name in src["fields"]:
                v = r.get(api_name)
                if v not in (None, ""):
                    row[our_name] = v
        record(src["name"], "OK", f"{len(recs)}행 수신 (최신 {max((x.get('report_date_as_yyyy_mm_dd') or '')[:10] for x in recs)})")

    if not got_any:
        return

    # 순포지션을 여기서 한 번만 계산해 둔다 — 회차마다 다시 빼지 않게.
    for row in merged.values():
        for lo, sh, net in (("noncomm_long", "noncomm_short", "noncomm_net"),
                            ("lev_long", "lev_short", "lev_net"),
                            ("am_long", "am_short", "am_net")):
            try:
                row[net] = int(row[lo]) - int(row[sh])
            except (KeyError, TypeError, ValueError):
                pass

    total, added, changed = merge_csv(
        os.path.join(OUT_DIR, "cftc_es.csv"), CFTC_FIELDS, list(merged.values()), "report_date")
    record("cftc_merge", "OK", f"cftc_es.csv 전체 {total}행", total, added, changed)


# ─────────────────────────────────────────────────────────────────────
# ② FINRA — 신용잔고 (월 1회)
# ─────────────────────────────────────────────────────────────────────
# 공식 배포 파일은 xlsx 하나이며 1997-01 부터 들어 있다.
# 백만 달러 단위로 적혀 있고, 열 이름이 해마다 조금씩 바뀌므로 «위치»가 아니라
# «머리글에 들어 있는 낱말»로 찾는다.
FINRA_URLS = [
    "https://www.finra.org/sites/default/files/2021-03/margin-statistics.xlsx",
    "https://www.finra.org/sites/default/files/margin-statistics.xlsx",
]
FINRA_FIELDS = ["month", "debit_balances_musd", "free_credit_cash_musd",
                "free_credit_margin_musd", "fetched_utc"]


def _norm_month(v):
    """'Aug-26' · '2026-08' · datetime 을 전부 '2026-08' 로 맞춘다."""
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m")
    s = str(v).strip()
    for fmt in ("%b-%y", "%b-%Y", "%Y-%m", "%m/%Y", "%B %Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m")
        except ValueError:
            continue
    return None


def fetch_finra():
    try:
        from openpyxl import load_workbook
    except ImportError:
        record("finra", "SKIP", "openpyxl 미설치 — 워크플로에 pip install openpyxl 이 필요하다")
        alarms.append("finra: openpyxl 미설치")
        return

    raw = None
    for url in FINRA_URLS:
        try:
            raw = http_get(url, accept="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            log(f"    FINRA 파일 수신: {url} ({len(raw)} B)")
            break
        except Exception as e:  # noqa: BLE001
            log(f"    {url} 실패: {e}")
    if raw is None:
        record("finra", "FAIL", "두 주소 모두 실패")
        alarms.append("finra: 파일을 받지 못했다")
        return

    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        grid = [list(r) for r in ws.iter_rows(values_only=True)]
    except Exception as e:  # noqa: BLE001
        record("finra", "FAIL", f"xlsx 해석 실패: {type(e).__name__} {e}")
        alarms.append(f"finra: xlsx 해석 실패 {e}")
        return

    # 머리글 행을 찾는다 — 'debit' 이 들어 있는 첫 행.
    hdr_i = None
    for i, row in enumerate(grid[:40]):
        joined = " ".join(str(c).lower() for c in row if c is not None)
        if "debit" in joined:
            hdr_i = i
            break
    if hdr_i is None:
        record("finra", "FAIL", "머리글 행('debit')을 찾지 못했다 — 파일 형식이 바뀌었다")
        alarms.append("finra: 머리글 행 없음")
        return

    hdr = [str(c).lower() if c is not None else "" for c in grid[hdr_i]]

    def find(*needles, avoid=()):
        for j, h in enumerate(hdr):
            if all(n in h for n in needles) and not any(a in h for a in avoid):
                return j
        return None

    c_debit = find("debit")
    c_cash = find("free credit", "cash") or find("cash account")
    c_marg = find("free credit", "margin") or find("margin account")
    if c_debit is None:
        record("finra", "FAIL", "차변잔액(debit) 열을 찾지 못했다")
        alarms.append("finra: debit 열 없음")
        return

    rows = []
    for row in grid[hdr_i + 1:]:
        if not row or row[0] is None:
            continue
        m = _norm_month(row[0])
        if not m:
            continue

        def cell(j):
            if j is None or j >= len(row) or row[j] is None:
                return ""
            v = row[j]
            return str(int(v)) if isinstance(v, (int, float)) else str(v).strip()

        if not cell(c_debit):
            continue
        rows.append({
            "month": m,
            "debit_balances_musd": cell(c_debit),
            "free_credit_cash_musd": cell(c_cash),
            "free_credit_margin_musd": cell(c_marg),
            "fetched_utc": RUN_UTC,
        })

    if not rows:
        record("finra", "EMPTY", "0행 — 쓰지 않는다")
        return

    total, added, changed = merge_csv(
        os.path.join(OUT_DIR, "finra_margin.csv"), FINRA_FIELDS, rows, "month")
    record("finra", "OK", f"최신 {max(r['month'] for r in rows)} · 전체 {total}행", total, added, changed)


# ─────────────────────────────────────────────────────────────────────
# ③ CBOE — 풋콜 비율 (일 1회)
# ─────────────────────────────────────────────────────────────────────
# 어디서 받나 [2026-09-24 · 90차 수정]
#    CBOE 공식 「일일 시장 통계」 웹페이지 한 곳이다.
#      https://www.cboe.com/us/options/market_statistics/daily/
#    내려받기 파일·API 는 찾지 못했다 — 공식 「과거 자료」 페이지의 풋콜 CSV 목록
#    (cdn.cboe.com/resources/options/volume_and_call_put_ratios/*.csv)은 페이지 문구상
#    2019-10-04 까지만 담고 갱신이 멈췄다. 처음 넣었던 cdn.cboe.com/api/global/us_indices/
#    daily_prices/*_pc.csv 세 주소는 «짐작»이었고 없는 파일이었다(2026-09-23 러너 시험 —
#    같은 서버의 VIX_History.csv 는 200 인데 이 주소만 AccessDenied 403). 그 주소들은 지웠다.
#    2026-09-24 러너 시험(수동 실행 · 07:18 ET): 브라우저 UA 로 HTTP 200 · 약 447 KB ·
#    화면 글자에서 TOTAL / INDEX / EQUITY PUT/CALL RATIO 세 값(0.86 · 1.04 · 0.46)이 뽑혔고,
#    9/23 의 전체·개별주식 값이 YCharts 와 일치했다(지수 값은 대조하지 않았다).
#
# 날짜를 어떻게 매기나 — 페이지 화면 글자에 날짜가 없다(위 시험 · dates=-).
#    그래서 «실행 시각»으로 매긴다. 기준은 미국 동부 시각(ET · 서머타임 자동 반영)이다.
#      ㉮ 평일 17:00 ET 이후 실행   → 그날   (장 마감 16:00 ET 뒤 1시간 이상 지남)
#      ㉯ 09:00 ET 이전 실행       → 직전 평일 (장 시작 전에는 직전 거래일 값을 보여 준다 — 위 시험)
#      ㉰ 토·일 실행                → 직전 금요일
#      ㉱ 평일 09:00~17:00 ET      → 저장하지 않는다 (장중·마감 직후 — 화면 값이 어느 날 것인지 모른다)
#    매일 22:30 UTC 실행(07:30 KST)은 18:30 EDT · 17:30 EST 라서 ㉮ 에 든다.
#    금요일 추가 실행(20:30 UTC = 16:30 EDT · 15:30 EST)은 ㉱ 라서 저장하지 않는다.
#    ⚠️ ㉮ 의 «17:00 ET 이후면 그날 마감 값을 보여 준다»는 «미시험»이다. 첫 며칠은 대화 창이
#       다른 자료처(YCharts 등)와 날짜별로 대조해 하루 밀림이 없는지 확인한다.
#    휴장일 — 목록을 두지 않는다. 휴장일에는 화면이 직전 거래일 값을 그대로 보여 줄 것이므로,
#       세 값이 이 파일의 마지막 행과 «셋 다 같으면» 저장하지 않고 NOTE 로 남긴다.
#       화면 갱신이 늦은 날도 같은 길로 걸러진다. 다음 날 아침(㉯) 실행이 있으면 그 행을 바로잡는다.
#
# 실패를 어떻게 알리나
#    페이지를 못 받거나, 세 값 중 하나라도 안 뽑히면 run_alarm.txt 에 남긴다(빨간불).
#    조용히 틀린 값을 쌓지 않는다 — 세 값이 다 뽑힌 실행만 한 행을 쓴다.
#
# 시험 한 줄(PROBE · 알람 없음) — ?dt= 로 지난 날짜를 부를 수 있는가
#    매 실행 ?dt=<매긴 날짜보다 3평일 앞> 으로 한 번 더 불러 세 값을 기록한다.
#    값이 오늘 화면과 다르면 지난 날짜를 부를 수 있다는 뜻이다(소급 수집이 가능해진다).
#    판정은 대화 창이 한다. 결론이 나면 이 시험은 지운다.
#
# 그동안에도 풋콜은 세션의 알파밴티지 도구(HISTORICAL_PUT_CALL_RATIO · SPY)가 따로 맡는다.
BROWSER_UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
              "(KHTML, like Gecko) Chrome/128.0 Safari/537.36")
CBOE_DAILY_PAGE = "https://www.cboe.com/us/options/market_statistics/daily/"
PC_LABELS = (("total", "TOTAL PUT/CALL RATIO"),
             ("index", "INDEX PUT/CALL RATIO"),
             ("equity", "EQUITY PUT/CALL RATIO"))
CBOE_FIELDS = ["date", "total_pc", "equity_pc", "index_pc", "fetched_utc"]
CBOE_CSV = os.path.join(OUT_DIR, "cboe_putcall.csv")
MARKET_TZ = "America/New_York"


def _cboe_page(url, timeout=60, retries=3):
    """브라우저 UA 로 페이지를 받아 글자로 돌려준다. 끝내 실패하면 예외를 올린다."""
    last = None
    for attempt in range(1, retries + 1):
        req = urllib.request.Request(url, headers={"User-Agent": BROWSER_UA, "Accept": "text/html"})
        try:
            with urllib.request.urlopen(req, timeout=timeout) as r:
                return r.read().decode("utf-8", errors="replace")
        except Exception as e:  # noqa: BLE001 — 어떤 실패든 재시도 대상
            last = e
            log(f"    시도 {attempt}/{retries} 실패: {e}")
            if attempt < retries:
                time.sleep(attempt * 5)
    raise last


def _pc_values(html):
    """화면 글자에서 세 비율을 뽑는다. 반환 — {"total": "0.86", ...} (못 뽑은 것은 빠진다)"""
    import re
    text = re.sub(r"<script[\s\S]*?</script>|<style[\s\S]*?</style>", " ", html, flags=re.I)
    text = re.sub(r"<[^>]+>", " ", text)
    text = " ".join(text.split())
    out = {}
    for key, label in PC_LABELS:
        m = re.search(re.escape(label) + r"\s*:?\s*([0-9]+\.[0-9]+)", text, flags=re.I)
        if m:
            out[key] = m.group(1)
    return out


def _date_strings(html, limit=5):
    """페이지 전체(스크립트 포함)의 날짜 모양 문자열 — 날짜 매김을 나중에 검산하려는 기록일 뿐이다."""
    import re
    found = re.findall(r"\b(20[0-9]{2}-[01][0-9]-[0-3][0-9]|[01][0-9]/[0-3][0-9]/20[0-9]{2})\b", html)
    return list(dict.fromkeys(found))[:limit]


def _label_date(now_utc=None):
    """실행 시각으로 날짜를 매긴다(위 ㉮~㉱). 반환 — (날짜 ISO 또는 None, 규칙 이름, ET 시각 문자열)"""
    from datetime import timedelta
    from zoneinfo import ZoneInfo
    et = (now_utc or datetime.now(timezone.utc)).astimezone(ZoneInfo(MARKET_TZ))
    d = et.date()
    stamp = et.strftime("%Y-%m-%d %H:%M ET")
    if d.weekday() >= 5:                       # 토(5)·일(6) → 직전 금요일
        return ((d - timedelta(days=d.weekday() - 4)).isoformat(), "weekend", stamp)
    if et.hour >= 17:
        return (d.isoformat(), "after_close", stamp)
    if et.hour < 9:
        return (_weekdays_before(d.isoformat(), 1), "before_open", stamp)
    return (None, "market_hours", stamp)


def _weekdays_before(iso, n):
    """iso 날짜보다 n 평일 앞의 날짜(ISO)."""
    from datetime import date, timedelta
    x = date.fromisoformat(iso)
    while n > 0:
        x -= timedelta(days=1)
        if x.weekday() < 5:
            n -= 1
    return x.isoformat()


def _last_row(path):
    if not os.path.exists(path):
        return None
    with open(path, "r", encoding="utf-8", newline="") as f:
        rows = list(csv.DictReader(f))
    return max(rows, key=lambda r: r["date"]) if rows else None


def fetch_cboe(now_utc=None):
    label, rule, stamp = _label_date(now_utc)
    try:
        html = _cboe_page(CBOE_DAILY_PAGE)
    except Exception as e:  # noqa: BLE001
        record("cboe_page", "FAIL", f"{type(e).__name__} {e}"[:600])
        alarms.append(f"cboe: 일일 통계 페이지를 받지 못했다 — {type(e).__name__} {e}")
        return

    vals = _pc_values(html)
    got = " ".join(f"{k}={vals.get(k, 'NA')}" for k, _ in PC_LABELS)
    ctx = (f"{stamp} · 규칙 {rule} · 매긴 날짜 {label or '-'} · "
           f"페이지 날짜 모양 {'|'.join(_date_strings(html)) or '-'} · bytes={len(html)}")
    missing = [k for k, _ in PC_LABELS if k not in vals]
    if missing:
        record("cboe_page", "FAIL", f"{got} · {ctx}"[:600])
        alarms.append(f"cboe: 페이지에서 {'·'.join(missing)} 비율을 뽑지 못했다 — 화면 문구가 바뀌었을 수 있다")
        return
    record("cboe_page", "OK", f"{got} · {ctx}"[:600])

    # 시험 한 줄 — 지난 날짜를 부를 수 있는가(알람 없음)
    probe_day = _weekdays_before(label or stamp[:10], 3)
    try:
        pv = _pc_values(_cboe_page(CBOE_DAILY_PAGE + "?dt=" + probe_day, retries=1))
        same = all(pv.get(k) == vals[k] for k, _ in PC_LABELS)
        record("cboe_probe_dt", "PROBE",
               f"dt={probe_day} · " + " ".join(f"{k}={pv.get(k, 'NA')}" for k, _ in PC_LABELS)
               + f" · 오늘 화면과 {'같음' if same else '다름'}")
    except Exception as e:  # noqa: BLE001
        record("cboe_probe_dt", "PROBE", f"dt={probe_day} · 실패 {type(e).__name__} {e}"[:600])

    if label is None:
        record("cboe_store", "NOTE", f"저장 안 함 — {stamp} 는 장중·마감 직후라 화면 값의 날짜를 정할 수 없다")
        return

    last = _last_row(CBOE_CSV)
    if last and last["date"] < label and all(str(last.get(f"{k}_pc", "")) == vals[k] for k, _ in PC_LABELS):
        record("cboe_store", "NOTE",
               f"저장 안 함 — 세 값이 마지막 행({last['date']})과 같다 · 휴장일이거나 화면 갱신 전으로 본다")
        return

    row = {"date": label, "fetched_utc": RUN_UTC}
    for k, _ in PC_LABELS:
        row[f"{k}_pc"] = vals[k]
    total, added, changed = merge_csv(CBOE_CSV, CBOE_FIELDS, [row], "date")
    record("cboe_merge", "OK", f"{label} 저장 · cboe_putcall.csv 전체 {total}행", total, added, changed)


def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    os.makedirs(META_DIR, exist_ok=True)
    if os.path.exists(ALARM_PATH):
        os.remove(ALARM_PATH)

    log(f"us_fetch 시작 {RUN_UTC}")
    for label, fn in (("CFTC 선물 포지션", fetch_cftc),
                      ("FINRA 신용잔고", fetch_finra),
                      ("CBOE 풋콜 비율", fetch_cboe)):
        log(f"── {label} ──")
        try:
            fn()
        except Exception as e:  # noqa: BLE001 — 한 소스가 죽어도 나머지는 간다
            record(label, "CRASH", f"{type(e).__name__} {e}")
            alarms.append(f"{label}: {type(e).__name__} {e}")

    # 소스 상태를 누적한다 — 「수집자도 자기 실패를 보고하지 않는다」(T-20)의 대책.
    write_header = not os.path.exists(HEALTH_PATH)
    with open(HEALTH_PATH, "a", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["run_utc", "source", "status", "rows",
                                          "added", "changed", "detail"])
        if write_header:
            w.writeheader()
        for r in health_rows:
            w.writerow(r)

    if alarms:
        with open(ALARM_PATH, "w", encoding="utf-8") as f:
            f.write(f"run_utc={RUN_UTC}\n")
            for a in alarms:
                f.write(a + "\n")
        log(f"알람 {len(alarms)}건 — {ALARM_PATH}")
        # 종료 코드 0 으로 끝낸다. 커밋 단계까지 가야 받은 것을 잃지 않는다.
        # 빨간불은 워크플로의 「수집 상태 확인」 단계가 run_alarm.txt 를 보고 켠다.
    log("us_fetch 끝")


if __name__ == "__main__":
    main()
    sys.exit(0)
